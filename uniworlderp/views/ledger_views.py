from .common_imports import *
from django.utils import timezone
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import io
from uniworlderp.models import ARInvoice, Payment, CustomerVendor, SalesEmployee


def get_ledger_filters(request):
    params = request.GET
    customer_id = params.get('customer', '')
    sales_employee_id = params.get('sales_employee', '')
    search = params.get('search', '').strip()
    txn_type = params.get('txn_type', '')
    preset = params.get('preset', 'day')
    start_date = params.get('start_date', '')
    end_date = params.get('end_date', '')

    today = timezone.localdate()
    if preset == 'month':
        start_date = today.replace(day=1).isoformat()
        end_date = today.isoformat()
    elif preset == 'year':
        start_date = today.replace(month=1, day=1).isoformat()
        end_date = today.isoformat()
    elif preset == 'range' and start_date and end_date:
        pass
    else:
        preset = 'day'
        start_date = end_date = today.isoformat()

    return {
        'customer_id': customer_id,
        'sales_employee_id': sales_employee_id,
        'search': search,
        'txn_type': txn_type,
        'preset': preset,
        'start_date': start_date,
        'end_date': end_date,
    }


def build_customer_ledger(customer, start_date, end_date, search='', txn_type=''):
    brought_forward = Decimal('0.00')

    invoices_before = ARInvoice.objects.filter(customer=customer)
    payments_before = Payment.objects.filter(customer=customer)
    if start_date:
        invoices_before = invoices_before.filter(invoice_date__lt=start_date)
        payments_before = payments_before.filter(payment_date__lt=start_date)
    else:
        invoices_before = invoices_before.none()
        payments_before = payments_before.none()

    invoiced_before = invoices_before.aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    paid_before = payments_before.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    brought_forward = invoiced_before - paid_before

    invoices = ARInvoice.objects.filter(customer=customer).select_related('sales_order')
    payments = Payment.objects.filter(customer=customer).select_related('invoice', 'method')
    if start_date:
        invoices = invoices.filter(invoice_date__gte=start_date)
        payments = payments.filter(payment_date__gte=start_date)
    if end_date:
        invoices = invoices.filter(invoice_date__lte=end_date)
        payments = payments.filter(payment_date__lte=end_date)

    search_lower = search.lower()

    rows = []
    for inv in invoices:
        ref = f"INV-{inv.id}" + (f" / SO-{inv.sales_order_id}" if inv.sales_order_id else "")
        matches_search = (
            not search_lower
            or search_lower in str(inv.id)
            or (inv.sales_order_id and search_lower in str(inv.sales_order_id))
            or search_lower in ref.lower()
        )
        matches_type = txn_type in ('', 'invoice')
        rows.append({
            'date': inv.invoice_date,
            'type': 'Invoice',
            'ref': ref,
            'debit': inv.total_amount,
            'credit': Decimal('0.00'),
            'notes': inv.notes or '',
            'link_url_name': 'customer_vendor:invoice_view',
            'link_pk': inv.id,
            'sort_key': (inv.invoice_date, 0, inv.id),
            'visible': matches_search and matches_type,
        })

    for pay in payments:
        method_label = pay.method.name if pay.method_id else ''
        type_label = f"Payment ({method_label})" if method_label else "Payment"
        if pay.received_by:
            type_label += f", rcvd by {pay.received_by}"
        ref = f"PAY-{pay.id}"
        matches_search = (
            not search_lower
            or search_lower in str(pay.id)
            or (pay.invoice_id and search_lower in str(pay.invoice_id))
            or search_lower in ref.lower()
        )
        matches_type = txn_type in ('', 'payment')
        rows.append({
            'date': pay.payment_date,
            'type': type_label,
            'ref': ref,
            'debit': Decimal('0.00'),
            'credit': pay.amount,
            'notes': pay.notes or '',
            'link_url_name': 'customer_vendor:payment_view',
            'link_pk': pay.id,
            'sort_key': (pay.payment_date, 1, pay.id),
            'visible': matches_search and matches_type,
        })

    rows.sort(key=lambda r: r['sort_key'])

    running_balance = brought_forward
    for row in rows:
        running_balance += row['debit'] - row['credit']
        row['balance'] = running_balance

    visible_rows = [row for row in rows if row['visible']]
    total_debit = sum((r['debit'] for r in visible_rows), Decimal('0.00'))
    total_credit = sum((r['credit'] for r in visible_rows), Decimal('0.00'))

    return {
        'brought_forward': brought_forward,
        'rows': visible_rows,
        'total_debit': total_debit,
        'total_credit': total_credit,
        'closing_balance': running_balance,
    }


class LedgerView(LoginRequiredMixin, View):
    template_name = 'ledger/ledger.html'

    def get(self, request):
        customers = CustomerVendor.objects.filter(entity_type='customer').order_by('name')
        sales_employees = SalesEmployee.objects.order_by('full_name')
        filters = get_ledger_filters(request)

        context = {
            'customers': customers,
            'sales_employees': sales_employees,
            **filters,
        }

        if filters['customer_id']:
            customer = get_object_or_404(CustomerVendor, pk=filters['customer_id'], entity_type='customer')
            ledger = build_customer_ledger(
                customer, filters['start_date'], filters['end_date'],
                search=filters['search'], txn_type=filters['txn_type'],
            )
            context['selected_customer'] = customer
            context['ledger'] = ledger
            context['invoiced_total'] = ledger['total_debit']
            context['collected_total'] = ledger['total_credit']
            context['outstanding_total'] = ledger['closing_balance']
            context['overdue_count'] = ARInvoice.objects.filter(
                customer=customer,
            ).exclude(payment_status='C').filter(
                due_date__lt=timezone.now()
            ).count()
        else:
            all_invoices = ARInvoice.objects.all()
            all_payments = Payment.objects.all()
            if filters['start_date']:
                all_invoices = all_invoices.filter(invoice_date__gte=filters['start_date'])
                all_payments = all_payments.filter(payment_date__gte=filters['start_date'])
            if filters['end_date']:
                all_invoices = all_invoices.filter(invoice_date__lte=filters['end_date'])
                all_payments = all_payments.filter(payment_date__lte=filters['end_date'])
            if filters['sales_employee_id']:
                all_invoices = all_invoices.filter(sales_employee_id=filters['sales_employee_id'])
                all_payments = all_payments.filter(sales_employee_id=filters['sales_employee_id'])

            context['invoiced_total'] = all_invoices.aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
            context['collected_total'] = all_payments.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
            context['outstanding_total'] = context['invoiced_total'] - context['collected_total']
            context['overdue_count'] = ARInvoice.objects.exclude(payment_status='C').filter(
                due_date__lt=timezone.now()
            ).count()

        return render(request, self.template_name, context)


class LedgerPrintView(LoginRequiredMixin, View):
    template_name = 'ledger/ledger_print.html'

    def get(self, request):
        customer_id = request.GET.get('customer', '')
        if not customer_id:
            return render(request, self.template_name, {'error': 'Select a customer to print a statement.'})

        customer = get_object_or_404(CustomerVendor, pk=customer_id, entity_type='customer')
        start_date = request.GET.get('start_date', '')
        end_date = request.GET.get('end_date', '')
        search = request.GET.get('search', '').strip()
        txn_type = request.GET.get('txn_type', '')
        ledger = build_customer_ledger(customer, start_date, end_date, search=search, txn_type=txn_type)

        context = {
            'customer': customer,
            'ledger': ledger,
            'start_date': start_date,
            'end_date': end_date,
            'user': request.user,
            'report_generated_at_formatted': timezone.localtime().strftime('%d/%m/%Y %I:%M %p'),
        }
        return render(request, self.template_name, context)


class LedgerExcelView(LoginRequiredMixin, View):
    def get(self, request):
        customer_id = request.GET.get('customer', '')
        customer = get_object_or_404(CustomerVendor, pk=customer_id, entity_type='customer')
        start_date = request.GET.get('start_date', '')
        end_date = request.GET.get('end_date', '')
        search = request.GET.get('search', '').strip()
        txn_type = request.GET.get('txn_type', '')
        ledger = build_customer_ledger(customer, start_date, end_date, search=search, txn_type=txn_type)

        wb = Workbook()
        ws = wb.active
        ws.title = "Customer Ledger"

        ws.merge_cells('A1:G1')
        title_cell = ws.cell(row=1, column=1, value=f"Customer Ledger - {customer.name}")
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center")

        headers = ['Date', 'Type', 'Ref', 'Debit (owed)', 'Credit (paid)', 'Balance', 'Notes']
        header_row = 3
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        row_num = header_row + 1
        ws.cell(row=row_num, column=1, value=start_date or 'Opening')
        ws.cell(row=row_num, column=2, value='Opening balance')
        ws.cell(row=row_num, column=6, value=float(ledger['brought_forward']))
        row_num += 1

        for row in ledger['rows']:
            ws.cell(row=row_num, column=1, value=row['date'].strftime('%d/%m/%Y') if row['date'] else '')
            ws.cell(row=row_num, column=2, value=row['type'])
            ws.cell(row=row_num, column=3, value=row['ref'])
            ws.cell(row=row_num, column=4, value=float(row['debit']))
            ws.cell(row=row_num, column=5, value=float(row['credit']))
            ws.cell(row=row_num, column=6, value=float(row['balance']))
            ws.cell(row=row_num, column=7, value=row['notes'])
            row_num += 1

        total_row = row_num + 1
        ws.cell(row=total_row, column=1, value='TOTALS:').font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=float(ledger['total_debit'])).font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=float(ledger['total_credit'])).font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=float(ledger['closing_balance'])).font = Font(bold=True)

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter if hasattr(column[0], 'column_letter') else None
            if column_letter:
                for cell in column:
                    try:
                        if hasattr(cell, 'value') and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename=ledger_{customer.name}_{start_date}_to_{end_date}.xlsx'
        return response
